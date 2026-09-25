"""Cell-map importer for Damodaran's fcffsimpleginzu.xlsx ("Input sheet") -> Assumptions.

The default map below follows the 2023-2025 layout of the Input sheet. Damodaran moves rows between versions,
so the map is overridable: pass `cell_map={...}` (key -> "B24") to `read_ginzu`, or store a full map under the
`ginzu_cell_map` pref. Values in the sheet are in the company's reporting units (millions usually) — the
resulting Assumptions carry them unchanged (bridge.debt/cash/shares/base_revenue/base_ebit).

Verify the map against your workbook: `read_ginzu(path)` returns `labels` (column A text next to each cell).
"""

from __future__ import annotations

import io
from pathlib import Path
from typing import Any, ClassVar

from app.valuation.assumptions import Assumptions

# key -> cell on the "Input sheet"
DEFAULT_CELL_MAP: dict[str, str] = {
    "revenue": "B8",  # Revenues (most recent 12 months)
    "ebit": "B9",  # Operating income or EBIT
    "interest_expense": "B10",
    "book_equity": "B11",
    "book_debt": "B12",
    "has_rd": "B13",  # Yes/No: capitalize R&D
    "has_leases": "B14",  # Yes/No: operating leases
    "cash": "B15",  # Cash and marketable securities
    "non_operating_assets": "B16",  # Cross holdings and other non-operating assets
    "minority": "B17",
    "shares": "B18",  # Number of shares outstanding
    "price": "B19",  # Current stock price
    "tax_effective": "B20",
    "tax_marginal": "B21",
    "growth_y1": "B24",  # Revenue growth rate for next year
    "margin_y1": "B25",  # Operating margin for next year
    "growth_2_5": "B26",  # Compounded annual revenue growth years 2-5
    "target_margin": "B27",  # Target pre-tax operating margin (year 10)
    "margin_convergence_year": "B28",
    "sales_to_capital_1_5": "B29",
    "sales_to_capital_6_10": "B30",
    "rf": "B33",  # Riskfree rate
    "wacc": "B34",  # Initial cost of capital
    "options_outstanding": "B38",
    "option_strike": "B39",
    "option_maturity": "B40",
    "stock_stdev": "B41",
    "terminal_roic_override": "B45",  # optional ROIC in stable growth
    "terminal_wacc_override": "B47",  # optional cost of capital in stable growth
    "terminal_g_override": "B50",  # optional growth rate in perpetuity (defaults to rf)
}
SHEET_NAMES = ("Input sheet", "Input Sheet", "Inputs", "Input")


def _cell_ref(ref: str) -> tuple[str, int]:
    col = "".join(ch for ch in ref if ch.isalpha())
    row = int("".join(ch for ch in ref if ch.isdigit()))
    return col, row


def read_ginzu(
    source: str | Path | bytes, cell_map: dict[str, str] | None = None, sheet: str | None = None
) -> dict[str, Any]:
    """Read mapped cells. Returns {"values": {key: value}, "labels": {key: column-A text}, "sheet": name, "cell_map": map}."""
    import openpyxl

    cm = {**DEFAULT_CELL_MAP, **(cell_map or {})}
    src = io.BytesIO(source) if isinstance(source, bytes) else str(source)
    wb = openpyxl.load_workbook(src, data_only=True)
    ws = None
    for name in ([sheet] if sheet else []) + list(SHEET_NAMES):
        if name and name in wb.sheetnames:
            ws = wb[name]
            break
    if ws is None:
        ws = wb[wb.sheetnames[0]]
    values: dict[str, Any] = {}
    labels: dict[str, Any] = {}
    for key, ref in cm.items():
        try:
            v = ws[ref].value
        except Exception:
            v = None
        values[key] = v
        _, row = _cell_ref(ref)
        labels[key] = ws[f"A{row}"].value
    return {"values": values, "labels": labels, "sheet": ws.title, "cell_map": cm}


def _f(v: Any) -> float | None:
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, int | float):
        return float(v)
    try:
        s = str(v).strip().replace(",", "")
        if s.endswith("%"):
            return float(s[:-1]) / 100
        return float(s)
    except ValueError:
        return None


def _rate(v: Any) -> float | None:
    x = _f(v)
    if x is None:
        return None
    return x / 100 if abs(x) > 1.0 else x


def _yes(v: Any) -> bool:
    return str(v).strip().lower() in ("yes", "y", "true", "1")


def ginzu_to_assumptions(values: dict[str, Any], forecast_years: int = 10) -> Assumptions:
    """Map the Input-sheet values onto Assumptions (10-year, sales-to-capital reinvestment, margin convergence)."""
    g1, g25 = _rate(values.get("growth_y1")), _rate(values.get("growth_2_5"))
    rf = _rate(values.get("rf"))
    g_t = _rate(values.get("terminal_g_override"))
    if g_t is None:
        g_t = rf
    path: list[float] | None = None
    if g1 is not None or g25 is not None:
        g1 = g1 if g1 is not None else g25
        g25 = g25 if g25 is not None else g1
        assert g1 is not None and g25 is not None
        path = [g1] + [g25] * 4
        gt = g_t if g_t is not None else g25
        # years 6-10: linear fade from g25 to terminal g (ginzu convention)
        path += [g25 + (gt - g25) * (i / 5) for i in range(1, 6)]
        path = path[:forecast_years]
    conv = _f(values.get("margin_convergence_year"))
    stc = _f(values.get("sales_to_capital_1_5"))
    a = Assumptions(
        scenario="base",
        forecast_years=forecast_years,
        base_revenue=_f(values.get("revenue")),
        base_ebit=_f(values.get("ebit")),
        revenue_growth_path=path,
        target_ebit_margin=_rate(values.get("target_margin")),
        margin_fade_years=int(conv) if conv else None,
        tax_rate=_rate(values.get("tax_effective")),
        tax_rate_marginal=_rate(values.get("tax_marginal")),
        reinvestment_method="sales_to_capital" if stc else None,
        sales_to_capital=stc,
        external_source="ginzu",
    )
    a.wacc.rf = rf
    a.wacc.wacc = _rate(values.get("wacc"))
    a.terminal.method = "gordon"
    a.terminal.g = g_t
    a.bridge.debt = _f(values.get("book_debt"))
    a.bridge.cash = _f(values.get("cash"))
    a.bridge.non_operating_assets = _f(values.get("non_operating_assets"))
    a.bridge.minority = _f(values.get("minority"))
    a.bridge.shares = _f(values.get("shares"))
    opts = _f(values.get("options_outstanding"))
    if opts:
        # simple treasury-stock dilution: options in the money at the average strike
        price, strike = _f(values.get("price")), _f(values.get("option_strike"))
        if price and strike is not None and price > strike:
            a.bridge.dilution = opts * (1 - strike / price)
    if _yes(values.get("has_leases")) is False:
        a.bridge.leases = 0.0
    a.note = "imported from fcffsimpleginzu Input sheet"
    return a


class GinzuImporter:
    """Importer facade used by the service (`kind='ginzu'`)."""

    kind: ClassVar[str] = "ginzu"

    def __init__(self, cell_map: dict[str, str] | None = None):
        self.cell_map = cell_map

    def load(self, source: str | Path | bytes) -> tuple[Assumptions, dict[str, Any]]:
        raw = read_ginzu(source, self.cell_map)
        return ginzu_to_assumptions(raw["values"]), raw
