"""External model: assumptions OR a ready fair value imported from CSV / XLSX / Google-Sheets CSV URL.

Template columns: key,value,unit,scenario,note. Keys mirror Assumptions paths (`wacc.beta`, `terminal.g`,
`revenue_growth`, `target_ebit_margin`, `bridge.shares` ...) or `fair_value` / `fair_value_low` / `fair_value_high`.
Unit `%` (or values like 8.5 for a rate) are converted to decimals. Empty scenario = base.
"""

from __future__ import annotations

import csv
import io
import re
from pathlib import Path
from typing import ClassVar

from pydantic import BaseModel, Field

from app.valuation.assumptions import Assumptions
from app.valuation.inputs import InputsSnapshot
from app.valuation.models.base import BaseValuationModel, ValuationResult
from app.valuation.models.fcff import FCFFModel

RATE_KEYS = {
    "revenue_growth",
    "target_ebit_margin",
    "tax_rate",
    "tax_rate_marginal",
    "capex_pct_revenue",
    "da_pct_revenue",
    "nwc_pct_revenue",
    "sbc_pct_revenue",
    "wacc.rf",
    "wacc.erp",
    "wacc.crp",
    "wacc.cost_of_debt",
    "wacc.weight_debt",
    "wacc.wacc",
    "terminal.g",
}
ALIASES = {
    "growth": "revenue_growth",
    "g": "terminal.g",
    "terminal_growth": "terminal.g",
    "wacc": "wacc.wacc",
    "beta": "wacc.beta",
    "rf": "wacc.rf",
    "erp": "wacc.erp",
    "crp": "wacc.crp",
    "cost_of_debt": "wacc.cost_of_debt",
    "margin": "target_ebit_margin",
    "ebit_margin": "target_ebit_margin",
    "shares": "bridge.shares",
    "net_debt": "bridge.net_debt",
    "cash": "bridge.cash",
    "debt": "bridge.debt",
    "minority": "bridge.minority",
    "leases": "bridge.leases",
    "exit_multiple": "terminal.multiple",
    "fv": "fair_value",
    "target_price": "fair_value",
}
FV_KEYS = {"fair_value", "fair_value_low", "fair_value_high"}
_SHEETS_RE = re.compile(r"docs\.google\.com/spreadsheets/d/([\w-]+)")
_URL_RE = re.compile(r"^https?://", re.I)


class ExternalImport(BaseModel):
    assumptions: dict[str, Assumptions] = Field(default_factory=dict)  # per scenario
    rows: list[dict] = Field(default_factory=list)
    errors: list[str] = Field(default_factory=list)
    source: str = ""

    @property
    def scenarios(self) -> list[str]:
        return list(self.assumptions)

    def base(self) -> Assumptions | None:
        return self.assumptions.get("base") or (
            next(iter(self.assumptions.values())) if self.assumptions else None
        )


def _num(v) -> float | None:
    if v is None:
        return None
    if isinstance(v, int | float):
        return float(v)
    s = str(v).strip().replace(",", "")
    if not s:
        return None
    pct = s.endswith("%")
    s = s.rstrip("%").strip()
    try:
        x = float(s)
    except ValueError:
        return None
    return x / 100 if pct else x


def parse_rows(rows: list[dict], source: str = "") -> ExternalImport:
    """rows: dicts with key,value[,unit,scenario,note] (header names case-insensitive)."""
    out = ExternalImport(source=source)
    per: dict[str, dict] = {}
    for raw in rows:
        r = {str(k).strip().lower(): v for k, v in raw.items() if k is not None}
        key = str(r.get("key") or "").strip()
        if not key or key.startswith("#"):
            continue
        key = ALIASES.get(key.lower(), key.lower())
        unit = str(r.get("unit") or "").strip().lower()
        scen = str(r.get("scenario") or "").strip().lower() or "base"
        val = r.get("value")
        if key in (
            "terminal.method",
            "reinvestment_method",
            "growth_fade",
            "scenario",
            "note",
            "terminal.multiple_metric",
            "multiples_use",
        ):
            per.setdefault(scen, {})[key] = str(val).strip() if val is not None else None
            continue
        x = _num(val)
        if x is None:
            out.errors.append(f"{key}: not a number ({val!r})")
            continue
        if unit in ("%", "pct", "percent") or (key in RATE_KEYS and abs(x) > 1.0):
            x = x / 100.0
        if unit in ("bn", "billion", "billions"):
            x *= 1e9
        elif unit in ("mm", "m", "million", "millions"):
            x *= 1e6
        elif unit in ("k", "thousand", "thousands"):
            x *= 1e3
        per.setdefault(scen, {})[key] = x
        out.rows.append({"key": key, "value": x, "unit": unit, "scenario": scen, "note": r.get("note")})
    for scen, flat in per.items():
        base = Assumptions(
            scenario=scen if scen in ("bear", "base", "bull") else "custom", external_source=source or None
        )
        for k, v in flat.items():
            if k in ("scenario", "note"):
                continue
            try:
                base = base.with_path(k, v)
            except Exception as e:
                out.errors.append(f"{k}: {e}"[:160])
        if scen not in ("bear", "base", "bull"):
            base.note = f"scenario {scen}"
        out.assumptions[scen] = base
    return out


def rows_from_csv(text: str) -> list[dict]:
    text = text.lstrip("﻿")
    sample = text[:2048]
    delim = ";" if sample.count(";") > sample.count(",") else ","
    reader = csv.DictReader(io.StringIO(text), delimiter=delim)
    if not reader.fieldnames or "key" not in [f.strip().lower() for f in reader.fieldnames]:
        # headerless "key,value" pairs
        rows = []
        for line in text.splitlines():
            parts = [p.strip() for p in line.split(delim)]
            if len(parts) >= 2 and parts[0]:
                rows.append(
                    {
                        "key": parts[0],
                        "value": parts[1],
                        "unit": parts[2] if len(parts) > 2 else "",
                        "scenario": parts[3] if len(parts) > 3 else "",
                    }
                )
        return rows
    return [dict(r) for r in reader]


def rows_from_xlsx(path_or_bytes: str | Path | bytes, sheet: str | None = None) -> list[dict]:
    import openpyxl

    src = io.BytesIO(path_or_bytes) if isinstance(path_or_bytes, bytes) else str(path_or_bytes)
    wb = openpyxl.load_workbook(src, data_only=True, read_only=True)
    ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = [str(h).strip().lower() if h is not None else "" for h in rows[0]]
    if "key" not in header:
        header = ["key", "value", "unit", "scenario", "note"][: len(rows[0])]
        body = rows
    else:
        body = rows[1:]
    out = []
    for r in body:
        d = {header[i]: r[i] for i in range(min(len(header), len(r))) if header[i]}
        if d.get("key"):
            out.append(d)
    return out


def sheets_csv_url(url: str) -> str:
    """Google Sheets share URL -> CSV export URL (keeps gid when present)."""
    m = _SHEETS_RE.search(url)
    if not m:
        return url
    gid = re.search(r"[#&?]gid=(\d+)", url)
    base = f"https://docs.google.com/spreadsheets/d/{m.group(1)}/export?format=csv"
    return base + (f"&gid={gid.group(1)}" if gid else "")


async def fetch_url_rows(url: str) -> list[dict]:
    from app.data.http import get_client
    from app.data.retry import raise_for_retry, retrying

    url = sheets_csv_url(url)
    client = get_client()
    async for attempt in retrying():
        with attempt:
            resp = await client.get(url)
            raise_for_retry(resp)
    ctype = resp.headers.get("content-type", "")
    if url.lower().endswith((".xlsx", ".xlsm")) or "spreadsheetml" in ctype:
        return rows_from_xlsx(resp.content)
    return rows_from_csv(resp.text)


def _read_source(source: str | Path | bytes, kind: str) -> list[dict]:
    """Blocking part of load_external (runs in a thread): file / bytes / inline CSV text -> rows."""
    if kind == "xlsx":
        return rows_from_xlsx(source)
    if isinstance(source, bytes):
        text = source.decode("utf-8", errors="replace")
    else:
        path = Path(str(source))
        text = path.read_text(encoding="utf-8") if path.exists() else str(source)
    return rows_from_csv(text)


async def load_external(
    source: str | Path | bytes, kind: str | None = None, filename: str | None = None
) -> ExternalImport:
    """kind: csv | xlsx | sheets | None (auto from extension/URL)."""
    import asyncio

    name = filename or (str(source) if not isinstance(source, bytes) else "")
    k = kind or (
        "sheets"
        if "docs.google.com" in name
        else "xlsx"
        if name.lower().endswith((".xlsx", ".xlsm"))
        else "csv"
    )
    if k == "sheets" or (isinstance(source, str) and _URL_RE.match(source)):
        rows = await fetch_url_rows(str(source))
    else:
        rows = await asyncio.to_thread(_read_source, source, k)
    return parse_rows(rows, source=name)


class ExternalModel(BaseValuationModel):
    id: ClassVar[str] = "external"
    name: ClassVar[str] = "External (imported fair value or assumptions)"
    version: ClassVar[str] = "1.0"
    inputs_required: ClassVar[set[str]] = set()
    reference: ClassVar[str | None] = "imported_fv"

    def run(self, inputs: InputsSnapshot, a: Assumptions, policies=None) -> ValuationResult:
        if a.fair_value is not None:
            price = inputs.market.price
            shares = inputs.shares()
            return ValuationResult(
                model_id=self.id,
                model_version=self.version,
                scenario=a.scenario or "base",
                value_per_share=a.fair_value,
                low=a.fair_value_low,
                high=a.fair_value_high,
                equity_value=a.fair_value * shares if shares else None,
                currency=inputs.currency,
                components={"fair_value": a.fair_value, "source": a.external_source},
                diagnostics={"price": price, "mode": "fair_value"},
            )
        if a.external_source is None and a.model_dump(exclude_none=True, exclude={"scenario"}) == {}:
            return ValuationResult(
                model_id=self.id,
                model_version=self.version,
                currency=inputs.currency,
                warnings=["nothing imported: no fair_value and no assumptions"],
            )
        res = FCFFModel().run(inputs, a, policies)
        res.model_id = self.id
        res.model_version = self.version
        res.diagnostics["mode"] = "assumptions -> fcff"
        res.diagnostics["source"] = a.external_source
        return res


MODEL = ExternalModel()
