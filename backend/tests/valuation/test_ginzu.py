import io

import openpyxl

from app.valuation.models.fcff import FCFFModel
from app.valuation.models.ginzu import DEFAULT_CELL_MAP, GinzuImporter, ginzu_to_assumptions, read_ginzu
from tests.valuation.conftest import synthetic_inputs


def _ginzu_bytes(cell_map=DEFAULT_CELL_MAP) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Input sheet"
    values = {
        "revenue": 1000,
        "ebit": 200,
        "interest_expense": 10,
        "book_equity": 800,
        "book_debt": 200,
        "has_rd": "No",
        "has_leases": "No",
        "cash": 100,
        "non_operating_assets": 0,
        "minority": 0,
        "shares": 100,
        "price": 20,
        "tax_effective": 0.25,
        "tax_marginal": 0.25,
        "growth_y1": 0.10,
        "margin_y1": 0.20,
        "growth_2_5": 0.08,
        "target_margin": 0.22,
        "margin_convergence_year": 5,
        "sales_to_capital_1_5": 2.0,
        "sales_to_capital_6_10": 2.0,
        "rf": 0.04,
        "wacc": 0.09,
    }
    for key, ref in cell_map.items():
        row = int("".join(c for c in ref if c.isdigit()))
        ws[f"A{row}"] = key.replace("_", " ")
        if key in values:
            ws[ref] = values[key]
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_ginzu_cell_map_import():
    raw = read_ginzu(_ginzu_bytes())
    assert raw["sheet"] == "Input sheet" and raw["values"]["revenue"] == 1000 and raw["labels"]["rf"] == "rf"
    a = ginzu_to_assumptions(raw["values"])
    assert a.forecast_years == 10 and a.base_revenue == 1000 and a.base_ebit == 200
    assert (
        a.revenue_growth_path[:5] == [0.10, 0.08, 0.08, 0.08, 0.08]
        and abs(a.revenue_growth_path[-1] - 0.04) < 1e-9
    )
    assert a.target_ebit_margin == 0.22 and a.margin_fade_years == 5 and a.sales_to_capital == 2.0
    assert a.wacc.wacc == 0.09 and a.wacc.rf == 0.04 and a.terminal.g == 0.04
    assert a.bridge.debt == 200 and a.bridge.cash == 100 and a.bridge.shares == 100 and a.bridge.leases == 0.0
    res = FCFFModel().run(synthetic_inputs(), a)
    assert res.value_per_share and res.value_per_share > 0
    assert res.diagnostics["resolved"]["reinvestment_method"] == "sales_to_capital"


def test_ginzu_cell_map_override():
    custom = {**DEFAULT_CELL_MAP, "rf": "B60", "wacc": "B61"}
    a, raw = GinzuImporter({"rf": "B60", "wacc": "B61"}).load(_ginzu_bytes(custom))
    assert raw["cell_map"]["rf"] == "B60" and a.wacc.rf == 0.04 and a.wacc.wacc == 0.09
